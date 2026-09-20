import csv
import json
import logging
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.parser.interface import (
    BaseParser,
    BlockType,
    ParserExecutionException,
    ParserRawResult,
    RawBlockData,
    RawPageData,
    RawTableData,
)

logger = logging.getLogger("jansetu.parser.excel")


class ExcelLayoutParser(BaseParser):
    """
    High-performance layout parser for government spreadsheets (.xlsx, .xls, .csv).
    Converts worksheets into canonical JanSetu structured document blocks and tables,
    enabling full extraction, normalization, and direct database publication.
    """

    @property
    def name(self) -> str:
        return "excel"

    @property
    def version(self) -> str:
        try:
            import openpyxl
            return f"openpyxl-{openpyxl.__version__}"
        except Exception:
            return "excel-1.0"

    def parse(
        self,
        file_path: Path,
        raw_output_dir: Path,
        max_pages: Optional[int] = None,
        timeout_seconds: Optional[int] = None,
        device: str = "auto",
        **kwargs: Any,
    ) -> ParserRawResult:
        start_time = time.perf_counter()
        raw_output_dir.mkdir(parents=True, exist_ok=True)

        if not file_path.exists():
            raise ParserExecutionException(f"Spreadsheet file not found at: {file_path}")

        ext = file_path.suffix.lower()
        raw_pages: List[RawPageData] = []
        raw_blocks_dump: List[Dict[str, Any]] = []

        try:
            if ext == ".xlsx":
                raw_pages, raw_blocks_dump = self._parse_xlsx(file_path, max_pages)
            elif ext == ".xls":
                raw_pages, raw_blocks_dump = self._parse_xls(file_path, max_pages)
            elif ext == ".csv":
                raw_pages, raw_blocks_dump = self._parse_csv(file_path)
            else:
                raise ParserExecutionException(f"Unsupported spreadsheet format: {ext}")

        except Exception as e:
            logger.error("Error during spreadsheet parsing for '%s': %s", file_path.name, e, exc_info=True)
            raise ParserExecutionException(f"Failed to parse spreadsheet document: {e}")

        total_pages = len(raw_pages)

        # Write raw parser dump to mineru_raw/raw_blocks.json (matches standard parser layout)
        raw_json_path = raw_output_dir / "raw_blocks.json"
        raw_dump_data = {
            "parser": self.name,
            "version": self.version,
            "page_count": total_pages,
            "pages": raw_blocks_dump,
        }
        with open(raw_json_path, "w", encoding="utf-8") as f:
            json.dump(raw_dump_data, f, ensure_ascii=False, indent=2)

        duration_ms = int((time.perf_counter() - start_time) * 1000)

        return ParserRawResult(
            success=True,
            parser_name=self.name,
            parser_version=self.version,
            page_count=total_pages,
            pages=raw_pages,
            raw_output_dir=raw_output_dir,
            raw_json=raw_dump_data,
            duration_ms=duration_ms,
        )

    def _parse_xlsx(self, file_path: Path, max_pages: Optional[int]) -> Tuple[List[RawPageData], List[Dict[str, Any]]]:
        import openpyxl

        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        if max_pages:
            sheet_names = sheet_names[:max_pages]

        raw_pages: List[RawPageData] = []
        raw_blocks_dump: List[Dict[str, Any]] = []

        for page_idx, sheet_name in enumerate(sheet_names):
            ws = wb[sheet_name]
            rows_data: List[List[str]] = []
            for row in ws.iter_rows(values_only=True):
                # Convert None to empty string and format cell values
                str_row = [str(c).strip() if c is not None else "" for c in row]
                # Keep if row has at least one non-empty value
                if any(str_row):
                    rows_data.append(str_row)

            page_num = page_idx + 1
            page_data, dump_data = self._process_sheet_data(page_num, sheet_name, rows_data)
            raw_pages.append(page_data)
            raw_blocks_dump.append(dump_data)

        wb.close()
        return raw_pages, raw_blocks_dump

    def _parse_xls(self, file_path: Path, max_pages: Optional[int]) -> Tuple[List[RawPageData], List[Dict[str, Any]]]:
        import xlrd

        wb = xlrd.open_workbook(str(file_path))
        sheet_count = wb.nsheets
        if max_pages:
            sheet_count = min(sheet_count, max_pages)

        raw_pages: List[RawPageData] = []
        raw_blocks_dump: List[Dict[str, Any]] = []

        for page_idx in range(sheet_count):
            sheet = wb.sheet_by_index(page_idx)
            rows_data: List[List[str]] = []
            for rx in range(sheet.nrows):
                str_row = [str(sheet.cell_value(rx, cx)).strip() for cx in range(sheet.ncols)]
                if any(str_row):
                    rows_data.append(str_row)

            page_num = page_idx + 1
            page_data, dump_data = self._process_sheet_data(page_num, sheet.name, rows_data)
            raw_pages.append(page_data)
            raw_blocks_dump.append(dump_data)

        return raw_pages, raw_blocks_dump

    def _parse_csv(self, file_path: Path) -> Tuple[List[RawPageData], List[Dict[str, Any]]]:
        encodings = ["utf-8-sig", "utf-8", "cp1252", "latin1"]
        rows_data: List[List[str]] = []

        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc, errors="replace") as f:
                    # Detect delimiter
                    sample = f.read(4096)
                    f.seek(0)
                    delimiter = ","
                    if "\t" in sample and sample.count("\t") > sample.count(","):
                        delimiter = "\t"
                    elif ";" in sample and sample.count(";") > sample.count(","):
                        delimiter = ";"

                    reader = csv.reader(f, delimiter=delimiter)
                    rows_data = [[str(c).strip() for c in row] for row in reader if any(c.strip() for c in row)]
                    break
            except Exception:
                continue

        page_data, dump_data = self._process_sheet_data(1, file_path.stem, rows_data)
        return [page_data], [dump_data]

    def _process_sheet_data(
        self,
        page_num: int,
        sheet_name: str,
        rows_data: List[List[str]],
    ) -> Tuple[RawPageData, Dict[str, Any]]:
        page_blocks: List[RawBlockData] = []
        order_index = 0

        # 1. Sheet Title Block
        title_block = RawBlockData(
            page_number=page_num,
            order_index=order_index,
            block_type=BlockType.TITLE if page_num == 1 else BlockType.HEADING,
            text=f"Worksheet: {sheet_name}",
            section_path=[sheet_name],
            bbox=[0, 0, 800, 30],
            metadata={"sheet_name": sheet_name},
        )
        page_blocks.append(title_block)
        order_index += 1

        if not rows_data:
            return (
                RawPageData(
                    page_number=page_num,
                    blocks=page_blocks,
                    raw_text=f"Worksheet {sheet_name} is empty.",
                    image_count=0,
                ),
                {
                    "page_number": page_num,
                    "block_count": len(page_blocks),
                    "blocks": [{"order_index": b.order_index, "block_type": b.block_type, "text": b.text} for b in page_blocks],
                },
            )

        # 2. Identify header row vs title rows
        # Often row 0 is the table header, unless it has only 1 non-empty cell while row 1 has multiple
        header_row_idx = 0
        if len(rows_data) > 1:
            first_row_non_empty = sum(1 for c in rows_data[0] if c)
            second_row_non_empty = sum(1 for c in rows_data[1] if c)
            if first_row_non_empty == 1 and second_row_non_empty > 1:
                # Row 0 is a preamble/document title
                preamble_text = " ".join(c for c in rows_data[0] if c)
                preamble_block = RawBlockData(
                    page_number=page_num,
                    order_index=order_index,
                    block_type=BlockType.HEADING,
                    text=preamble_text,
                    section_path=[sheet_name, preamble_text[:50]],
                    bbox=[0, 30, 800, 60],
                    metadata={},
                )
                page_blocks.append(preamble_block)
                order_index += 1
                header_row_idx = 1

        # Extract headers
        raw_headers = rows_data[header_row_idx]
        # Normalize headers (fill blank headers as Column_N)
        headers: List[str] = []
        for col_i, h in enumerate(raw_headers):
            col_name = h.strip() if h else f"Column_{col_i + 1}"
            headers.append(col_name)

        # Pad or slice data rows to match header length
        data_rows: List[List[str]] = []
        for r in rows_data[header_row_idx + 1:]:
            padded = r[:len(headers)] + [""] * max(0, len(headers) - len(r))
            data_rows.append(padded)

        # 3. Create Markdown Table
        table_lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * len(headers)) + " |",
        ]
        for row in data_rows:
            escaped_row = [c.replace("|", "\\|").replace("\n", " ") for c in row]
            table_lines.append("| " + " | ".join(escaped_row) + " |")

        table_markdown = "\n".join(table_lines)

        table_data = RawTableData(
            headers=headers,
            rows=data_rows,
            markdown=table_markdown,
        )

        table_block = RawBlockData(
            page_number=page_num,
            order_index=order_index,
            block_type=BlockType.TABLE,
            text=f"Table ({sheet_name}): {', '.join(headers[:6])}",
            section_path=[sheet_name, "Schemes Data"],
            table_data=table_data,
            bbox=[0, 70, 800, 300],
            metadata={"row_count": len(data_rows), "col_count": len(headers), "sheet": sheet_name},
        )
        page_blocks.append(table_block)
        order_index += 1

        # 4. Generate Semantic Row Paragraph Blocks for Entity Extraction
        # For each row, synthesize a rich descriptive paragraph linking headers to values.
        # This guarantees LLM and heuristic extraction agents capture every scheme field!
        row_descriptions: List[str] = []
        for row_idx, row_values in enumerate(data_rows):
            field_pairs = []
            for h, v in zip(headers, row_values):
                if v:
                    field_pairs.append(f"{h}: {v}")
            if field_pairs:
                row_text = f"Scheme Entry (Row {row_idx + 1}): " + " | ".join(field_pairs)
                row_descriptions.append(row_text)

                row_block = RawBlockData(
                    page_number=page_num,
                    order_index=order_index,
                    block_type=BlockType.PARAGRAPH,
                    text=row_text,
                    section_path=[sheet_name, f"Entry {row_idx + 1}"],
                    bbox=[0, 310 + (row_idx * 25), 800, 330 + (row_idx * 25)],
                    metadata={"row_index": row_idx + 1, "sheet": sheet_name},
                )
                page_blocks.append(row_block)
                order_index += 1

        # Build comprehensive raw_text for the page
        full_page_text_parts = [
            f"# {sheet_name}",
            "",
            "## Data Table",
            table_markdown,
            "",
            "## Structured Scheme Records",
            "\n\n".join(row_descriptions),
        ]
        full_page_text = "\n".join(full_page_text_parts)

        # Prepare dump for raw_blocks.json
        dump_entry = {
            "page_number": page_num,
            "block_count": len(page_blocks),
            "sheet_name": sheet_name,
            "blocks": [
                {
                    "order_index": blk.order_index,
                    "block_type": blk.block_type,
                    "text": blk.text,
                    "section_path": blk.section_path,
                    "bbox": blk.bbox,
                    "has_table": blk.table_data is not None,
                }
                for blk in page_blocks
            ],
        }

        return (
            RawPageData(
                page_number=page_num,
                blocks=page_blocks,
                raw_text=full_page_text,
                image_count=0,
            ),
            dump_entry,
        )
